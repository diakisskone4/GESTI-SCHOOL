"""Liste des enseignants au format Excel (.xlsx), même présentation que la liste des élèves."""
import io
import re

from django.utils import timezone
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font
from openpyxl.utils import get_column_letter

from apps.staff.models import Enseignant
from apps.students.exports import _BORDURE, _ENTETE, _VIOLET, _ZEBRE

COLONNES = [
    ("N°", 6), ("Matricule", 20), ("Nom", 18), ("Prénom", 22), ("Sexe", 10), ("Téléphone", 16),
    ("Email", 26), ("Grade", 22), ("Diplôme", 20), ("Spécialité", 20), ("Matières enseignées", 34),
    ("Contrat", 12), ("Date d'embauche", 15), ("Statut", 10), ("Ajouté le", 14),
]


def classeur_enseignants(etablissement):
    enseignants = (
        Enseignant.objects.filter(etablissement=etablissement)
        .select_related("user").prefetch_related("matieres_enseignees").order_by("nom", "prenom")
    )
    wb = Workbook()
    ws = wb.active
    ws.title = "Enseignants"
    derniere = get_column_letter(len(COLONNES))
    ws["A1"] = etablissement.nom
    ws["A1"].font = Font(bold=True, size=14, color=_VIOLET)
    ws["A2"] = f"Liste des enseignants — {len(enseignants)} enseignant(s) — éditée le {timezone.localtime():%d/%m/%Y à %H:%M}"
    ws["A2"].font = Font(italic=True, color="64748B")
    ws.merge_cells(f"A1:{derniere}1")
    ws.merge_cells(f"A2:{derniere}2")
    ws.append([])
    ws.append([c for c, _ in COLONNES])
    for cellule in ws[4]:
        cellule.font = Font(bold=True, color="FFFFFF")
        cellule.fill = _ENTETE
        cellule.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        cellule.border = _BORDURE
    ws.row_dimensions[4].height = 28

    for numero, e in enumerate(enseignants, start=1):
        ws.append([
            numero, e.matricule, e.nom.upper(), e.prenom, e.get_sexe_display(), e.telephone,
            e.email_personnel or (e.user.email if e.user else ""), e.grade_academique, e.diplome, e.specialite,
            ", ".join(m.nom for m in e.matieres_enseignees.all()), e.type_contrat, e.date_embauche,
            "Actif" if e.actif else "Inactif", timezone.localtime(e.created_at).date(),
        ])
        for cellule in ws[ws.max_row]:
            cellule.border = _BORDURE
            if numero % 2 == 0:
                cellule.fill = _ZEBRE
        for colonne in (13, 15):
            ws.cell(row=ws.max_row, column=colonne).number_format = "DD/MM/YYYY"

    for index, (_, largeur) in enumerate(COLONNES, start=1):
        ws.column_dimensions[get_column_letter(index)].width = largeur
    ws.freeze_panes = "A5"
    if enseignants:
        ws.auto_filter.ref = f"A4:{derniere}{ws.max_row}"
    ws.page_setup.orientation = "landscape"
    ws.page_setup.fitToWidth = 1
    ws.print_title_rows = "4:4"
    tampon = io.BytesIO()
    wb.save(tampon)
    return tampon.getvalue()


def nom_fichier_enseignants(etablissement):
    sigle = re.sub(r"[^A-Za-z0-9]+", "-", etablissement.sigle).strip("-")
    return f"enseignants_{sigle}_{timezone.localdate():%Y-%m-%d}.xlsx"
