from django.test import TestCase

# Create your tests here.


class ListeEnseignantsTests(TestCase):
    def setUp(self):
        import datetime

        from django.utils import timezone
        from rest_framework.test import APIClient

        from apps.accounts.models import User
        from apps.core.tests import creer_etablissement
        from apps.staff.models import Enseignant

        self.etab, _, _ = creer_etablissement("ENS")
        for i in range(230):
            e = Enseignant.objects.create(etablissement=self.etab, nom=f"NOM{i:03d}", prenom="Prof", telephone=f"7000{i:04d}")
            Enseignant.objects.filter(pk=e.pk).update(created_at=timezone.now() - datetime.timedelta(minutes=230 - i))
        admin = User.objects.create_user(email="a@ens.ml", password="x", role="admin", etablissement_courant=self.etab)
        admin.etablissements.add(self.etab)
        self.client = APIClient(HTTP_HOST="localhost")
        self.client.force_authenticate(admin)

    def test_les_200_derniers_ajoutes_d_abord(self):
        r = self.client.get("/api/v1/enseignants/")
        self.assertEqual(r.data["count"], 230)
        self.assertEqual(len(r.data["results"]), 200)
        self.assertEqual(r.data["results"][0]["nom"], "NOM229")  # le plus récent en premier
        self.assertEqual(r.data["results"][-1]["nom"], "NOM030")
        alpha = self.client.get("/api/v1/enseignants/", {"ordering": "nom", "page_size": 25}).data["results"]
        self.assertEqual(alpha[0]["nom"], "NOM000")

    def test_export_excel(self):
        import io

        from openpyxl import load_workbook

        r = self.client.get("/api/v1/enseignants/export_excel/")
        self.assertEqual(r.status_code, 200)
        ws = load_workbook(io.BytesIO(r.content)).active
        self.assertEqual(ws["B4"].value, "Matricule")
        self.assertEqual(ws.max_row - 4, 230)
