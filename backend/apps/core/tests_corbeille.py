"""Tests de la corbeille (restauration des éléments supprimés) et de l'archive complète de l'établissement."""
import datetime
import io
import json
import zipfile

from django.test import TestCase, override_settings
from django.utils import timezone
from openpyxl import load_workbook
from rest_framework.test import APIClient

from apps.accounts.models import User
from apps.core.models import Classe, ElementSupprime, Matiere
from apps.core.tests import creer_etablissement
from apps.staff.models import Enseignant
from apps.students.models import Absence, ArchiveEleves, Eleve, Inscription


class CorbeilleTests(TestCase):
    def setUp(self):
        self.etab, self.classe, self.eleve = creer_etablissement("CSM")
        self.annee = self.classe.annee_scolaire
        self.inscription = Inscription.objects.create(eleve=self.eleve, classe=self.classe, annee_scolaire=self.annee)
        Absence.objects.create(inscription=self.inscription, date=datetime.date(2026, 10, 5), motif="Malade")
        self.admin = User.objects.create_user(email="admin@csm.ml", password="x", role="admin", first_name="Awa",
                                              last_name="Traore", etablissement_courant=self.etab)
        self.admin.etablissements.add(self.etab)
        self.client = APIClient(HTTP_HOST="localhost")
        self.client.force_authenticate(self.admin)

    def corbeille(self):
        return self.client.get("/api/v1/corbeille/").data["results"]

    def test_eleve_supprime_puis_restaure_avec_ses_donnees_liees(self):
        eleve_id = self.eleve.id
        self.assertEqual(self.client.delete(f"/api/v1/eleves/{eleve_id}/").status_code, 204)
        self.assertFalse(Eleve.objects.filter(pk=eleve_id).exists())
        self.assertFalse(Absence.objects.exists())

        [element] = self.corbeille()
        self.assertEqual(element["type_objet"], "Élève")
        self.assertEqual(element["nb_objets"], 3)  # élève + inscription + absence
        self.assertEqual(element["supprime_par_nom"], "Awa Traore")

        r = self.client.post(f"/api/v1/corbeille/{element['id']}/restaurer/")
        self.assertEqual(r.status_code, 200, r.data)
        self.assertIn("2 élément(s) lié(s)", r.data["detail"])
        eleve = Eleve.objects.get(pk=eleve_id)
        self.assertEqual(eleve.matricule, self.eleve.matricule)
        self.assertEqual(eleve.inscriptions.get().classe, self.classe)
        self.assertEqual(Absence.objects.get().motif, "Malade")
        self.assertEqual(self.corbeille(), [])

    def test_reference_remise_a_vide_puis_retablie(self):
        enseignant = Enseignant.objects.create(etablissement=self.etab, nom="KONE", prenom="Issa")
        matiere = Matiere.objects.create(etablissement=self.etab, nom="Maths")
        enseignant.matieres_enseignees.add(matiere)
        Classe.objects.filter(pk=self.classe.pk).update(professeur_principal=enseignant)

        self.assertEqual(self.client.delete(f"/api/v1/enseignants/{enseignant.pk}/").status_code, 204)
        self.assertIsNone(Classe.objects.get(pk=self.classe.pk).professeur_principal_id)

        [element] = self.corbeille()
        self.assertEqual(self.client.post(f"/api/v1/corbeille/{element['id']}/restaurer/").status_code, 200)
        restaure = Enseignant.objects.get(pk=enseignant.pk)  # héritage Employe -> Enseignant reconstitué
        self.assertEqual(restaure.nom, "KONE")
        self.assertEqual(list(restaure.matieres_enseignees.all()), [matiere])
        self.assertEqual(Classe.objects.get(pk=self.classe.pk).professeur_principal_id, enseignant.pk)

    def test_restauration_impossible_si_le_parent_a_disparu(self):
        self.client.delete(f"/api/v1/inscriptions/{self.inscription.pk}/")
        self.client.delete(f"/api/v1/classes/{self.classe.pk}/")
        inscription = next(e for e in self.corbeille() if e["type_objet"] == "Inscription")
        r = self.client.post(f"/api/v1/corbeille/{inscription['id']}/restaurer/")
        self.assertEqual(r.status_code, 400)
        self.assertIn("Restaurez-le d'abord", r.data["detail"])
        self.assertFalse(Inscription.objects.exists())  # rien n'a été modifié
        # Dans le bon ordre, tout revient
        classe = next(e for e in self.corbeille() if e["type_objet"] == "Classe")
        self.assertEqual(self.client.post(f"/api/v1/corbeille/{classe['id']}/restaurer/").status_code, 200)
        self.assertEqual(self.client.post(f"/api/v1/corbeille/{inscription['id']}/restaurer/").status_code, 200)
        self.assertEqual(Inscription.objects.get().classe_id, self.classe.pk)

    def test_compte_utilisateur_restaure_avec_ses_etablissements(self):
        prof = User.objects.create_user(email="prof@csm.ml", password="x", role="enseignant", etablissement_courant=self.etab)
        prof.etablissements.add(self.etab)
        self.assertEqual(self.client.delete(f"/api/v1/utilisateurs/{prof.pk}/").status_code, 204)
        [element] = self.corbeille()
        self.assertEqual(self.client.post(f"/api/v1/corbeille/{element['id']}/restaurer/").status_code, 200)
        prof = User.objects.get(email="prof@csm.ml")
        self.assertEqual(list(prof.etablissements.all()), [self.etab])
        self.assertTrue(prof.check_password("x"))

    def test_cloisonnement_suppression_definitive_et_vidage(self):
        self.client.delete(f"/api/v1/eleves/{self.eleve.pk}/")
        autre, _, eleve_b = creer_etablissement("AUT")
        admin_b = User.objects.create_user(email="b@aut.ml", password="x", role="admin", etablissement_courant=autre)
        admin_b.etablissements.add(autre)
        client_b = APIClient(HTTP_HOST="localhost")
        client_b.force_authenticate(admin_b)
        self.assertEqual(client_b.get("/api/v1/corbeille/").data["count"], 0)
        element_a = ElementSupprime.objects.get()
        self.assertEqual(client_b.post(f"/api/v1/corbeille/{element_a.pk}/restaurer/").status_code, 404)
        self.assertEqual(client_b.delete(f"/api/v1/corbeille/{element_a.pk}/").status_code, 404)

        client_b.delete(f"/api/v1/eleves/{eleve_b.pk}/")
        self.assertEqual(self.client.post("/api/v1/corbeille/vider/").status_code, 200)
        self.assertEqual(list(ElementSupprime.objects.values_list("etablissement_id", flat=True)), [autre.pk])

    @override_settings(CORBEILLE_DUREE_JOURS=30)
    def test_elements_trop_anciens_effaces(self):
        self.client.delete(f"/api/v1/eleves/{self.eleve.pk}/")
        ElementSupprime.objects.update(supprime_le=timezone.now() - datetime.timedelta(days=31))
        self.assertEqual(self.corbeille(), [])

    def test_reserve_aux_administrateurs(self):
        prof = User.objects.create_user(email="p@csm.ml", password="x", role="enseignant", etablissement_courant=self.etab)
        prof.etablissements.add(self.etab)
        client = APIClient(HTTP_HOST="localhost")
        client.force_authenticate(prof)
        self.assertEqual(client.get("/api/v1/corbeille/").status_code, 403)


class ArchiveCompleteTests(TestCase):
    def setUp(self):
        self.etab, self.classe, self.eleve = creer_etablissement("CSM")
        Inscription.objects.create(eleve=self.eleve, classe=self.classe, annee_scolaire=self.classe.annee_scolaire)
        Enseignant.objects.create(etablissement=self.etab, nom="KONE", prenom="Issa")
        self.autre, _, _ = creer_etablissement("AUT")
        self.admin = User.objects.create_user(email="admin@csm.ml", password="secret-123", role="admin",
                                              etablissement_courant=self.etab)
        self.admin.etablissements.add(self.etab)
        self.client = APIClient(HTTP_HOST="localhost")
        self.client.force_authenticate(self.admin)

    def test_archive_complete_de_l_etablissement(self):
        r = self.client.post("/api/v1/archives-eleves/creer/", {"portee": "etablissement"}, format="json")
        self.assertEqual(r.status_code, 201, r.data)
        self.assertEqual(r.data["portee"], "etablissement")
        zf = zipfile.ZipFile(io.BytesIO(self.client.get(f"/api/v1/archives-eleves/{r.data['id']}/telecharger/").content))
        noms = zf.namelist()
        self.assertIn("donnees/students.eleve.json", noms)
        self.assertIn("donnees/staff.enseignant.json", noms)
        self.assertIn("donnees/core.classe.json", noms)
        self.assertNotIn("donnees/plateforme.abonnement.json", noms)
        utilisateurs = json.loads(zf.read("donnees/accounts.user.json"))
        self.assertEqual([u["fields"]["email"] for u in utilisateurs], ["admin@csm.ml"])
        self.assertNotIn("password", utilisateurs[0]["fields"])
        # Uniquement les données de CET établissement
        eleves = json.loads(zf.read("donnees/students.eleve.json"))
        self.assertEqual([e["pk"] for e in eleves], [str(self.eleve.pk)])
        classeur = next(n for n in noms if "_donnees_" in n)
        wb = load_workbook(io.BytesIO(zf.read(classeur)))
        self.assertIn("Résumé", wb.sheetnames)
        self.assertIn("Élèves", wb.sheetnames)
        # La liste des archives distingue les deux sortes
        self.assertEqual(self.client.get("/api/v1/archives-eleves/", {"portee": "etablissement"}).data["count"], 1)
        self.assertEqual(self.client.get("/api/v1/archives-eleves/", {"portee": "eleves"}).data["count"], 0)

    @override_settings(ARCHIVES_ETABLISSEMENT_MAX=2)
    def test_retention_des_archives_completes(self):
        for _ in range(3):
            self.client.post("/api/v1/archives-eleves/creer/", {"portee": "etablissement"}, format="json")
        self.assertEqual(ArchiveEleves.objects.filter(portee="etablissement").count(), 2)
