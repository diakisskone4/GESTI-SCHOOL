"""Liste des élèves au format Excel (.xlsx), par classe.

Utilisée par le bouton « Exporter Excel » de la page Élèves et par les archives ZIP (apps/students/archives.py).
"""
import io
import re
from collections import OrderedDict

from django.db.models import Prefetch
from django.utils import timezone
from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

from apps.students.models import Eleve, Inscription

COLONNES = [
    ("N°", 6), ("Matricule", 22), ("Nom", 18), ("Prénom", 22), ("Sexe", 10), ("Date de naissance", 16),
    ("Lieu de naissance", 18), ("Nationalité", 14), ("Nom du père", 22), ("Nom de la mère", 22),
    ("Contact d'urgence", 24), ("Téléphone", 16), ("Adresse", 26), ("Statut", 10),
]
SANS_CLASSE = "Sans classe"

_VIOLET = "5F34DD"
_ENTETE = PatternFill("solid", fgColor=_VIOLET)
_ZEBRE = PatternFill("solid", fgColor="F6F5FB")
_BORDURE = Border(*(Side(style="thin", color="E2E8F0"),) * 4)


def eleves_avec_classe(etablissement, classe=None):
    """Élèves de l'établissement avec leur inscription active (la plus récente), triés par nom.
    Avec `classe` : uniquement les élèves inscrits (inscription active) dans cette classe."""
    actives = Inscription.objects.filter(statut="active").select_related("classe__niveau", "annee_scolaire").order_by("-date_inscription")
    qs = Eleve.objects.filter(etablissement=etablissement).prefetch_related(
        Prefetch("inscriptions", queryset=actives, to_attr="inscriptions_actives"),
    ).order_by("nom", "prenom")
    if classe is not None:
        qs = qs.filter(inscriptions__classe=classe, inscriptions__statut="active").distinct()
    for eleve in qs:
        eleve.inscription_active = eleve.inscriptions_actives[0] if eleve.inscriptions_actives else None
        yield eleve


def _ligne(numero, eleve):
    return [
        numero, eleve.matricule, eleve.nom.upper(), eleve.prenom, eleve.get_sexe_display(),
        eleve.date_naissance, eleve.lieu_naissance, eleve.nationalite, eleve.nom_pere, eleve.nom_mere,
        eleve.contact_urgence_nom, eleve.contact_urgence_telephone, eleve.adresse, "Actif" if eleve.actif else "Inactif",
    ]


def _nom_feuille(nom, deja_pris):
    """Nom d'onglet Excel valide (31 caractères, sans []:*?/\\) et unique dans le classeur."""
    base = re.sub(r"[\[\]:*?/\\]", "-", nom).strip()[:31] or "Classe"
    candidat, i = base, 2
    while candidat.lower() in deja_pris:
        suffixe = f" ({i})"
        candidat, i = base[: 31 - len(suffixe)] + suffixe, i + 1
    deja_pris.add(candidat.lower())
    return candidat


def _remplir_feuille(ws, etablissement, titre, eleves, avec_classe=False):
    colonnes = list(COLONNES)
    if avec_classe:
        colonnes.insert(4, ("Classe", 18))
    derniere = get_column_letter(len(colonnes))

    ws["A1"] = etablissement.nom
    ws["A1"].font = Font(bold=True, size=14, color=_VIOLET)
    ws["A2"] = f"{titre} — {len(eleves)} élève(s) — édité le {timezone.localtime():%d/%m/%Y à %H:%M}"
    ws["A2"].font = Font(italic=True, color="64748B")
    ws.merge_cells(f"A1:{derniere}1")
    ws.merge_cells(f"A2:{derniere}2")

    ws.append([])
    ws.append([c for c, _ in colonnes])
    for cellule in ws[4]:
        cellule.font = Font(bold=True, color="FFFFFF")
        cellule.fill = _ENTETE
        cellule.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        cellule.border = _BORDURE
    ws.row_dimensions[4].height = 28

    for numero, eleve in enumerate(eleves, start=1):
        valeurs = _ligne(numero, eleve)
        if avec_classe:
            insc = eleve.inscription_active
            valeurs.insert(4, insc.classe.nom if insc else SANS_CLASSE)
        ws.append(valeurs)
        for cellule in ws[ws.max_row]:
            cellule.border = _BORDURE
            if numero % 2 == 0:
                cellule.fill = _ZEBRE
        ws.cell(row=ws.max_row, column=6 + avec_classe).number_format = "DD/MM/YYYY"

    for index, (_, largeur) in enumerate(colonnes, start=1):
        ws.column_dimensions[get_column_letter(index)].width = largeur
    ws.freeze_panes = "A5"
    if eleves:
        ws.auto_filter.ref = f"A4:{derniere}{ws.max_row}"
    ws.page_setup.orientation = "landscape"
    ws.page_setup.fitToWidth = 1
    ws.print_title_rows = "4:4"


def classeur_eleves(etablissement, classe=None):
    """Contenu .xlsx de la liste des élèves.

    - Avec `classe` : un seul onglet, les élèves de cette classe.
    - Sans : un onglet « Tous les élèves » puis un onglet par classe (et « Sans classe » si besoin).
    """
    wb = Workbook()
    ws = wb.active
    if classe is not None:
        ws.title = _nom_feuille(classe.nom, set())
        libelle = f"Classe {classe.nom} ({classe.annee_scolaire.libelle})"
        _remplir_feuille(ws, etablissement, libelle, list(eleves_avec_classe(etablissement, classe)))
    else:
        eleves = list(eleves_avec_classe(etablissement))
        ws.title = "Tous les élèves"
        _remplir_feuille(ws, etablissement, "Tous les élèves", eleves, avec_classe=True)
        par_classe = OrderedDict()
        for eleve in sorted(eleves, key=lambda e: (
            e.inscription_active is None,
            e.inscription_active.classe.niveau.ordre if e.inscription_active else 0,
            e.inscription_active.classe.nom if e.inscription_active else "",
        )):
            nom = eleve.inscription_active.classe.nom if eleve.inscription_active else SANS_CLASSE
            par_classe.setdefault(nom, []).append(eleve)
        noms_pris = {ws.title.lower()}
        for nom, liste in par_classe.items():
            feuille = wb.create_sheet(_nom_feuille(nom, noms_pris))
            titre = nom if nom == SANS_CLASSE else f"Classe {nom}"
            _remplir_feuille(feuille, etablissement, titre, sorted(liste, key=lambda e: (e.nom, e.prenom)))

    tampon = io.BytesIO()
    wb.save(tampon)
    return tampon.getvalue()


def nom_fichier_excel(etablissement, classe=None):
    cible = re.sub(r"[^A-Za-z0-9]+", "-", classe.nom).strip("-") if classe is not None else "toutes-classes"
    return f"eleves_{etablissement.sigle}_{cible}_{timezone.localdate():%Y-%m-%d}.xlsx"
