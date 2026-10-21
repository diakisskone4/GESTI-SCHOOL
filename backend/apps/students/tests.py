"""Tests de la page Élèves : pagination et filtre par classe, export Excel, archives ZIP."""
import datetime
import io
import json
import zipfile

from django.test import TestCase, override_settings
from openpyxl import load_workbook
from rest_framework.test import APIClient

from apps.accounts.models import User
from apps.core.models import Classe, Niveau
from apps.core.tests import creer_etablissement
from apps.students.models import ArchiveEleves, Eleve, Inscription


class ElevesTestsBase(TestCase):
    def setUp(self):
        self.etab, self.classe_a, self.eleve_existant = creer_etablissement("CSM")
        niveau = Niveau.objects.create(etablissement=self.etab, nom="5e", cycle="secondaire_1", ordre=2)
        self.classe_b = Classe.objects.create(
            etablissement=self.etab, annee_scolaire=self.classe_a.annee_scolaire, niveau=niveau, nom="5e B",
        )
        Inscription.objects.create(eleve=self.eleve_existant, classe=self.classe_a, annee_scolaire=self.classe_a.annee_scolaire)
        for i in range(30):
            eleve = Eleve.objects.create(etablissement=self.etab, nom=f"NOM{i:02d}", prenom="Awa", sexe="F",
                                         date_naissance=datetime.date(2012, 1, 1) + datetime.timedelta(days=i))
            classe = self.classe_a if i < 20 else self.classe_b
            Inscription.objects.create(eleve=eleve, classe=classe, annee_scolaire=classe.annee_scolaire)
        Eleve.objects.create(etablissement=self.etab, nom="SANSCLASSE", prenom="Ali", sexe="M", date_naissance=datetime.date(2011, 5, 5))

        self.admin = User.objects.create_user(email="admin@csm.ml", password="x", role="admin", first_name="Souleymane",
                                              last_name="Dembele", etablissement_courant=self.etab)
        self.admin.etablissements.add(self.etab)
        self.client = APIClient(HTTP_HOST="localhost")
        self.client.force_authenticate(self.admin)


class ListeElevesTests(ElevesTestsBase):
    def test_compteur_total_et_taille_de_page(self):
        r = self.client.get("/api/v1/eleves/")
        self.assertEqual(r.data["count"], 32)  # le total, pas seulement les 25 de la page
        self.assertEqual(len(r.data["results"]), 25)
        r = self.client.get("/api/v1/eleves/", {"page_size": 100})
        self.assertEqual(len(r.data["results"]), 32)

    def test_filtre_par_classe_cote_serveur(self):
        r = self.client.get("/api/v1/eleves/", {"classe": self.classe_a.id, "page_size": 100})
        self.assertEqual(r.data["count"], 21)
        self.assertTrue(all(e["classe_actuelle"] == self.classe_a.nom for e in r.data["results"]))
        self.assertEqual(self.client.get("/api/v1/eleves/", {"classe": self.classe_b.id}).data["count"], 10)
        sans = self.client.get("/api/v1/eleves/", {"classe": "aucune"}).data
        self.assertEqual([e["nom"] for e in sans["results"]], ["SANSCLASSE"])
        self.assertEqual(self.client.get("/api/v1/eleves/", {"classe": "x"}).status_code, 400)


class ExportExcelTests(ElevesTestsBase):
    def classeur(self, **params):
        r = self.client.get("/api/v1/eleves/export_excel/", params)
        self.assertEqual(r.status_code, 200)
        self.assertIn(".xlsx", r["Content-Disposition"])
        return load_workbook(io.BytesIO(r.content))

    def test_export_d_une_classe(self):
        wb = self.classeur(classe=self.classe_b.id)
        self.assertEqual(wb.sheetnames, ["5e B"])
        ws = wb.active
        self.assertEqual(ws["A1"].value, self.etab.nom)
        self.assertEqual(ws["B4"].value, "Matricule")
        noms = [ws.cell(row=r, column=3).value for r in range(5, ws.max_row + 1)]
        self.assertEqual(len(noms), 10)
        self.assertEqual(noms, sorted(noms))

    def test_export_toutes_les_classes_un_onglet_par_classe(self):
        wb = self.classeur()
        self.assertEqual(wb.sheetnames, ["Tous les élèves", "6e CSM", "5e B", "Sans classe"])
        self.assertEqual(wb["Tous les élèves"].max_row - 4, 32)
        self.assertEqual(wb["6e CSM"].max_row - 4, 21)

    def test_classe_d_un_autre_etablissement_refusee(self):
        _, classe_autre, _ = creer_etablissement("AUT")
        r = self.client.get("/api/v1/eleves/export_excel/", {"classe": classe_autre.id})
        self.assertEqual(r.status_code, 404)

    def test_reserve_aux_administrateurs(self):
        prof = User.objects.create_user(email="prof@csm.ml", password="x", role="enseignant", etablissement_courant=self.etab)
        prof.etablissements.add(self.etab)
        client = APIClient(HTTP_HOST="localhost")
        client.force_authenticate(prof)
        self.assertEqual(client.get("/api/v1/eleves/export_excel/").status_code, 403)


class ArchivesTests(ElevesTestsBase):
    def lire_zip(self, archive_id):
        r = self.client.get(f"/api/v1/archives-eleves/{archive_id}/telecharger/")
        self.assertEqual(r.status_code, 200)
        self.assertEqual(r["Content-Type"], "application/zip")
        return zipfile.ZipFile(io.BytesIO(r.content))

    def test_archive_creee_et_regroupee_apres_une_inscription(self):
        with self.captureOnCommitCallbacks(execute=True):
            r = self.client.post("/api/v1/eleves/", {
                "nom": "NOUVEAU", "prenom": "Moussa", "sexe": "M", "date_naissance": "2013-03-03",
            }, format="json")
        self.assertEqual(r.status_code, 201, r.data)
        with self.captureOnCommitCallbacks(execute=True):
            self.client.post("/api/v1/inscriptions/", {
                "eleve": r.data["id"], "classe": self.classe_b.id, "annee_scolaire": self.classe_b.annee_scolaire_id,
            }, format="json")
        # Les deux opérations rapprochées tiennent dans une seule archive
        self.assertEqual(ArchiveEleves.objects.count(), 1)
        archive = self.client.get("/api/v1/archives-eleves/").data["results"][0]
        self.assertEqual([o["type"] for o in archive["operations"]], ["creation", "inscription"])
        self.assertEqual(archive["nb_eleves"], 33)
        self.assertEqual(archive["auteur_nom"], "Souleymane Dembele")
        zf = self.lire_zip(archive["id"])
        noms = zf.namelist()
        self.assertTrue(any(n.endswith(".xlsx") for n in noms))
        self.assertIn("donnees.json", noms)
        donnees = json.loads(zf.read("donnees.json"))
        nouveau = next(e for e in donnees["eleves"] if e["nom"] == "NOUVEAU")
        self.assertEqual(nouveau["inscriptions"][0]["classe"], "5e B")

    def test_suppression_garde_l_etat_avant(self):
        eleve = Eleve.objects.get(nom="NOM05")
        with self.captureOnCommitCallbacks(execute=True):
            self.assertEqual(self.client.delete(f"/api/v1/eleves/{eleve.id}/").status_code, 204)
        archive = ArchiveEleves.objects.get()
        operations = json.loads(self.lire_zip(archive.id).read("operations.json"))
        self.assertEqual(operations[0]["type"], "suppression")
        self.assertEqual(operations[0]["avant"]["nom"], "NOM05")
        self.assertEqual(operations[0]["avant"]["inscriptions"][0]["classe"], self.classe_a.nom)
        self.assertEqual(archive.nb_eleves, 31)

    def test_operation_refusee_aucune_archive(self):
        _, classe_autre, _ = creer_etablissement("AUT")
        with self.captureOnCommitCallbacks(execute=True):
            r = self.client.post("/api/v1/inscriptions/", {
                "eleve": str(self.eleve_existant.id), "classe": classe_autre.id,
                "annee_scolaire": classe_autre.annee_scolaire_id,
            }, format="json")
        self.assertEqual(r.status_code, 403)
        self.assertFalse(ArchiveEleves.objects.exists())

    @override_settings(ARCHIVES_ELEVES_MAX=3)
    def test_archive_manuelle_et_retention(self):
        for _ in range(5):
            self.assertEqual(self.client.post("/api/v1/archives-eleves/creer/").status_code, 201)
        self.assertEqual(ArchiveEleves.objects.filter(etablissement=self.etab).count(), 3)

    def test_archives_cloisonnees_par_etablissement(self):
        self.client.post("/api/v1/archives-eleves/creer/")
        autre_etab, _, _ = creer_etablissement("AUT")
        admin_b = User.objects.create_user(email="b@aut.ml", password="x", role="admin", etablissement_courant=autre_etab)
        admin_b.etablissements.add(autre_etab)
        client_b = APIClient(HTTP_HOST="localhost")
        client_b.force_authenticate(admin_b)
        self.assertEqual(client_b.get("/api/v1/archives-eleves/").data["count"], 0)
        archive_a = ArchiveEleves.objects.get()
        self.assertEqual(client_b.get(f"/api/v1/archives-eleves/{archive_a.id}/telecharger/").status_code, 404)
