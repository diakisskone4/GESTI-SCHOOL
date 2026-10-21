"""Archivage ZIP automatique des données élèves.

Après chaque opération sur les élèves ou leurs inscriptions (création, modification, changement de classe,
suppression), une archive ZIP de l'établissement est enregistrée :

    eleves_<SIGLE>_<date>.xlsx   la liste complète, un onglet par classe (apps/students/exports.py)
    donnees.json                 toutes les données élèves et inscriptions, lisibles par un programme
    operations.json              les opérations couvertes : qui, quoi, quand, et l'état AVANT une suppression
    LISEZMOI.txt                 description du contenu

Les opérations rapprochées d'un même utilisateur (moins de 2 minutes) sont regroupées dans la même archive :
inscrire un élève puis l'affecter à une classe ne produit qu'une archive. Seules les N dernières archives
de chaque établissement sont conservées (réglage ARCHIVES_ELEVES_MAX, 50 par défaut).

L'archivage ne doit jamais faire échouer l'opération elle-même : en cas d'erreur, elle est seulement journalisée.
"""
import datetime
import io
import json
import logging
import zipfile

from django.conf import settings
from django.core.serializers.json import DjangoJSONEncoder
from django.db import transaction
from django.utils import timezone

from apps.students.exports import classeur_eleves, eleves_avec_classe
from apps.students.models import ArchiveEleves

logger = logging.getLogger(__name__)

FENETRE_REGROUPEMENT = datetime.timedelta(minutes=2)

CHAMPS_ELEVE = [
    "matricule", "nom", "prenom", "sexe", "date_naissance", "lieu_naissance", "nationalite", "nom_pere",
    "nom_mere", "adresse", "groupe_sanguin", "allergies_ou_besoins_speciaux", "contact_urgence_nom",
    "contact_urgence_telephone", "ecole_provenance", "actif",
]


def donnees_eleve(eleve):
    """Instantané lisible d'un élève et de ses inscriptions (utilisé aussi avant une suppression)."""
    return {
        "id": str(eleve.pk),
        **{champ: getattr(eleve, champ) for champ in CHAMPS_ELEVE},
        "inscriptions": [
            {
                "annee_scolaire": i.annee_scolaire.libelle, "classe": i.classe.nom, "statut": i.statut,
                "type": i.type_inscription, "redoublant": i.redoublant, "boursier": i.boursier,
                "date_inscription": i.date_inscription,
            }
            for i in eleve.inscriptions.select_related("classe", "annee_scolaire").order_by("-date_inscription")
        ],
    }


def _limite():
    return getattr(settings, "ARCHIVES_ELEVES_MAX", 50)


def construire_zip(etablissement, operations):
    eleves = list(eleves_avec_classe(etablissement))
    horodatage = timezone.localtime()
    nom_excel = f"eleves_{etablissement.sigle}_{horodatage:%Y-%m-%d_%H%M}.xlsx"
    donnees = {
        "etablissement": {"id": etablissement.pk, "nom": etablissement.nom, "sigle": etablissement.sigle},
        "genere_le": horodatage,
        "nb_eleves": len(eleves),
        "eleves": [donnees_eleve(e) for e in eleves],
    }
    lisezmoi = (
        f"Archive des données élèves — {etablissement.nom}\n"
        f"Générée le {horodatage:%d/%m/%Y à %H:%M:%S} ({len(eleves)} élève(s)).\n\n"
        f"{nom_excel} : liste des élèves, un onglet par classe (ouvrir avec Excel ou LibreOffice).\n"
        "donnees.json : toutes les données élèves et inscriptions.\n"
        "operations.json : opérations à l'origine de cette archive (auteur, date, état avant suppression).\n"
    )
    tampon = io.BytesIO()
    with zipfile.ZipFile(tampon, "w", compression=zipfile.ZIP_DEFLATED) as zf:
        zf.writestr(nom_excel, classeur_eleves(etablissement))
        zf.writestr("donnees.json", json.dumps(donnees, cls=DjangoJSONEncoder, ensure_ascii=False, indent=2))
        zf.writestr("operations.json", json.dumps(operations, cls=DjangoJSONEncoder, ensure_ascii=False, indent=2))
        zf.writestr("LISEZMOI.txt", lisezmoi)
    return tampon.getvalue(), len(eleves), f"archive_{etablissement.sigle}_{horodatage:%Y-%m-%d_%H%M%S}.zip"


def creer_archive(etablissement, auteur, type_operation, description, avant=None, regrouper=True):
    """Crée (ou complète, si une archive récente du même auteur existe) l'archive ZIP de l'établissement."""
    operation = {
        "date": timezone.now(), "type": type_operation, "description": description,
        "auteur": auteur.get_full_name() if auteur else None,
    }
    if avant is not None:
        operation["avant"] = avant
    operation = json.loads(json.dumps(operation, cls=DjangoJSONEncoder))

    recente = None
    if regrouper and auteur is not None:
        recente = ArchiveEleves.objects.filter(
            etablissement=etablissement, portee="eleves", auteur=auteur, modifie_le__gte=timezone.now() - FENETRE_REGROUPEMENT,
        ).order_by("-modifie_le").first()
    operations = (recente.operations if recente else []) + [operation]
    contenu, nb_eleves, nom = construire_zip(etablissement, operations)

    archive = recente or ArchiveEleves(etablissement=etablissement, auteur=auteur, nom_fichier=nom)
    archive.operations = operations
    archive.contenu = contenu
    archive.taille = len(contenu)
    archive.nb_eleves = nb_eleves
    archive.save()

    # Rétention : seules les N dernières archives de l'établissement sont gardées
    anciennes = ArchiveEleves.objects.filter(etablissement=etablissement, portee="eleves").order_by("-cree_le").values_list("pk", flat=True)[_limite():]
    ArchiveEleves.objects.filter(pk__in=list(anciennes)).delete()
    return archive


def archiver_apres_operation(etablissement_id, auteur, type_operation, description, avant=None):
    """À appeler pendant une opération : l'archive est créée une fois l'opération validée en base
    (rien n'est archivé si elle est annulée), et une erreur d'archivage ne la fait jamais échouer."""
    if not etablissement_id:
        return

    def _archiver():
        from apps.core.models import Etablissement

        try:
            etablissement = Etablissement.objects.get(pk=etablissement_id)
            creer_archive(etablissement, auteur, type_operation, description, avant=avant)
        except Exception:  # noqa: BLE001 - l'archivage est un filet de sécurité, jamais bloquant
            logger.exception("Échec de l'archivage des élèves (établissement %s)", etablissement_id)

    transaction.on_commit(_archiver)


# ---------------------------------------------------------------------------
# Archive complète de l'établissement
# ---------------------------------------------------------------------------
CHAMPS_UTILISATEUR = ["email", "first_name", "last_name", "role", "telephone", "is_active", "date_joined",
                      "date_derniere_activite", "etablissement_courant"]  # jamais le mot de passe


def donnees_etablissement(etablissement):
    """Toutes les données qui appartiennent à l'établissement, par modèle : ce que la base supprimerait
    avec lui (même mécanisme que la corbeille, sans rien supprimer), plus ses comptes utilisateurs."""
    from django.db.models.deletion import Collector

    from apps.accounts.models import User
    from apps.core.corbeille import MODELES_EXCLUS

    collector = Collector(using="default", origin=etablissement)
    collector.collect([etablissement])
    par_modele = {}
    for model, instances in collector.data.items():
        par_modele.setdefault(model, []).extend(instances)
    for qs in collector.fast_deletes:
        par_modele.setdefault(qs.model, []).extend(qs)
    par_modele = {
        m: sorted(objs, key=lambda o: str(o.pk)) for m, objs in par_modele.items()
        if objs and m._meta.label_lower not in MODELES_EXCLUS and not m._meta.auto_created
    }
    utilisateurs = list(User.objects.filter(etablissements=etablissement).order_by("last_name", "first_name"))
    return par_modele, utilisateurs


def _valeur_excel(valeur):
    import decimal
    import uuid

    if isinstance(valeur, datetime.datetime):
        return timezone.localtime(valeur).replace(tzinfo=None) if timezone.is_aware(valeur) else valeur
    if isinstance(valeur, (str, int, float, bool, datetime.date, decimal.Decimal)) or valeur is None:
        return valeur
    if isinstance(valeur, uuid.UUID):
        return str(valeur)
    if isinstance(valeur, (dict, list)):
        return json.dumps(valeur, ensure_ascii=False, cls=DjangoJSONEncoder)
    return str(valeur)


def classeur_etablissement(etablissement, par_modele, utilisateurs):
    """Un onglet par type de données (classes, élèves, notes, factures...) avec toutes les colonnes."""
    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill
    from openpyxl.utils import get_column_letter

    from apps.accounts.models import User
    from apps.students.exports import _nom_feuille

    wb = Workbook()
    resume = wb.active
    resume.title = "Résumé"
    resume.append([etablissement.nom])
    resume["A1"].font = Font(bold=True, size=14, color="5F34DD")
    resume.append([f"Archive complète du {timezone.localtime():%d/%m/%Y à %H:%M}"])
    resume.append([])
    resume.append(["Données", "Nombre"])
    noms_pris = {"résumé"}
    colonnes_utilisateur = [f for f in User._meta.concrete_fields if f.name in CHAMPS_UTILISATEUR]
    tables = [("Utilisateurs", colonnes_utilisateur, utilisateurs)] + [
        (str(m._meta.verbose_name_plural).capitalize(),
         [f for f in m._meta.concrete_fields if f.get_internal_type() != "BinaryField"], objs)
        for m, objs in par_modele.items()
    ]
    def sans_accents(texte):
        import unicodedata

        return unicodedata.normalize("NFKD", texte).encode("ascii", "ignore").decode().lower()

    for titre, colonnes, objets in sorted(tables, key=lambda t: sans_accents(t[0])):
        resume.append([titre, len(objets)])
        ws = wb.create_sheet(_nom_feuille(titre, noms_pris))
        ws.append([str(f.verbose_name).capitalize() for f in colonnes])
        for cellule in ws[1]:
            cellule.font = Font(bold=True, color="FFFFFF")
            cellule.fill = PatternFill("solid", fgColor="5F34DD")
        for obj in objets:
            ws.append([_valeur_excel(getattr(obj, f.attname)) for f in colonnes])
        ws.freeze_panes = "A2"
        for i in range(1, len(colonnes) + 1):
            ws.column_dimensions[get_column_letter(i)].width = 18
    resume.column_dimensions["A"].width = 40
    tampon = io.BytesIO()
    wb.save(tampon)
    return tampon.getvalue()


def construire_zip_etablissement(etablissement):
    from django.core import serializers

    par_modele, utilisateurs = donnees_etablissement(etablissement)
    horodatage = timezone.localtime()
    nb_eleves = sum(len(objs) for m, objs in par_modele.items() if m._meta.label_lower == "students.eleve")
    contenu = {m._meta.label_lower: len(objs) for m, objs in par_modele.items()}
    contenu["accounts.user"] = len(utilisateurs)
    manifeste = {
        "etablissement": {"id": etablissement.pk, "nom": etablissement.nom, "sigle": etablissement.sigle},
        "genere_le": horodatage,
        "contenu": contenu,
    }
    lisezmoi = "\n".join([
        f"Archive complète — {etablissement.nom}",
        f"Générée le {horodatage:%d/%m/%Y à %H:%M:%S}.",
        "",
        "*_donnees_*.xlsx : toutes les données, un onglet par type (classes, élèves, notes, factures, paie...).",
        "*_eleves_par_classe_*.xlsx : liste des élèves, un onglet par classe.",
        "donnees/*.json : les mêmes données au format JSON de Django (réimportables par l'équipe technique).",
        "manifeste.json : nombre d'éléments par type.",
        "Les mots de passe des comptes ne sont jamais inclus.",
    ])
    tampon = io.BytesIO()
    with zipfile.ZipFile(tampon, "w", compression=zipfile.ZIP_DEFLATED) as zf:
        zf.writestr(f"{etablissement.sigle}_donnees_{horodatage:%Y-%m-%d}.xlsx",
                    classeur_etablissement(etablissement, par_modele, utilisateurs))
        zf.writestr(f"{etablissement.sigle}_eleves_par_classe_{horodatage:%Y-%m-%d}.xlsx", classeur_eleves(etablissement))
        for modele, objets in par_modele.items():
            zf.writestr(f"donnees/{modele._meta.label_lower}.json", serializers.serialize("json", objets, indent=2))
        zf.writestr("donnees/accounts.user.json",
                    serializers.serialize("json", utilisateurs, fields=CHAMPS_UTILISATEUR, indent=2))
        zf.writestr("manifeste.json", json.dumps(manifeste, cls=DjangoJSONEncoder, ensure_ascii=False, indent=2))
        zf.writestr("LISEZMOI.txt", lisezmoi)
    return tampon.getvalue(), nb_eleves, f"archive_complete_{etablissement.sigle}_{horodatage:%Y-%m-%d_%H%M%S}.zip"


def creer_archive_etablissement(etablissement, auteur):
    contenu, nb_eleves, nom = construire_zip_etablissement(etablissement)
    archive = ArchiveEleves.objects.create(
        etablissement=etablissement, portee="etablissement", auteur=auteur, nom_fichier=nom,
        contenu=contenu, taille=len(contenu), nb_eleves=nb_eleves,
        operations=[{"date": timezone.now().isoformat(), "type": "etablissement",
                     "description": "Archive complète de l'établissement",
                     "auteur": auteur.get_full_name() if auteur else None}],
    )
    limite = getattr(settings, "ARCHIVES_ETABLISSEMENT_MAX", 10)
    anciennes = ArchiveEleves.objects.filter(etablissement=etablissement, portee="etablissement") \
        .order_by("-cree_le").values_list("pk", flat=True)[limite:]
    ArchiveEleves.objects.filter(pk__in=list(anciennes)).delete()
    return archive
